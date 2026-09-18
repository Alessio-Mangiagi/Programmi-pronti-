"""
Import dell'albero WBS di un cantiere da Excel (.xlsx) o CSV.

Formato accettato (intestazioni riconosciute in italiano/inglese, senza distinzione
di maiuscole; se manca l'intestazione: prima colonna codice, seconda nome):

    codice | nome                | padre    (opzionale)  | livello (opzionale)
    01     | Opere strutturali   |                       | 1
    01.02  | Solai               | 01                    | 2

La gerarchia si ricava, in ordine di precedenza:
1. colonna "padre" (codice della voce superiore);
2. colonna "livello" (1 = radice; ogni riga è figlia dell'ultima riga di livello inferiore);
3. codice puntato ("01.02" è figlia di "01", "01.02.03" di "01.02"; separatori . - /);
4. altrimenti tutte le voci sono radici.

Le voci si aggiornano per codice (stesso codice nello stesso cantiere = stessa voce):
un secondo import con il file corretto rinomina/risposta invece di duplicare.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from typing import Optional

HEADERS = {
    "code": {"codice", "cod", "code", "wbs", "id", "codice wbs", "wbs code"},
    "name": {"nome", "name", "descrizione", "description", "titolo", "title", "voce", "denominazione", "attivita", "attività"},
    "parent": {"padre", "parent", "parent_code", "codice padre", "superiore", "genitore", "parent code"},
    "level": {"livello", "level", "liv", "lvl"},
}
CODE_SEPARATORS = re.compile(r"[.\-/]")
MAX_ROWS = 5000


@dataclass
class ImportRow:
    row: int                      # numero di riga nel file (1-based, come in Excel)
    code: Optional[str]
    name: str
    parent_code: Optional[str]    # codice del padre risolto (None = radice)
    error: Optional[str] = None


@dataclass
class ParsedFile:
    rows: list[ImportRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)   # errori di file (non di riga)


def _norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    out = []
    for r in ws.iter_rows(values_only=True):
        out.append([_norm(c) for c in r])
        if len(out) > MAX_ROWS:
            break
    wb.close()
    return out


def _read_csv(data: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), dialect)
    out = []
    for r in reader:
        out.append([_norm(c) for c in r])
        if len(out) > MAX_ROWS:
            break
    return out


def read_table(data: bytes, filename: str = "") -> list[list[str]]:
    """Righe grezze (stringhe) da xlsx o csv; il tipo si riconosce dai byte (zip = xlsx) o dall'estensione."""
    if data[:2] == b"PK" or filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            return _read_xlsx(data)
        except Exception:
            raise ValueError("cannot read xlsx")
    if filename.lower().endswith((".xls",)):
        raise ValueError("xls not supported: save as xlsx or csv")
    return _read_csv(data)


def _detect_columns(header: list[str]) -> Optional[dict[str, int]]:
    """Indice di colonna per campo, se la prima riga è un'intestazione riconoscibile."""
    cols: dict[str, int] = {}
    for i, cell in enumerate(header):
        key = cell.strip().lower()
        for fieldname, aliases in HEADERS.items():
            if key in aliases and fieldname not in cols:
                cols[fieldname] = i
    return cols if "name" in cols or "code" in cols else None


def parse(data: bytes, filename: str = "") -> ParsedFile:
    table = read_table(data, filename)
    table = [r for r in table if any(c for c in r)]
    out = ParsedFile()
    if not table:
        out.errors.append("empty file")
        return out
    if len(table) > MAX_ROWS:
        out.errors.append(f"too many rows (max {MAX_ROWS})")
        return out

    cols = _detect_columns(table[0])
    if cols is None:
        cols = {"code": 0, "name": 1} if len(table[0]) > 1 else {"name": 0}
        body = enumerate(table, start=1)
    else:
        body = enumerate(table[1:], start=2)

    def cell(r: list[str], key: str) -> str:
        i = cols.get(key)
        return r[i] if i is not None and i < len(r) else ""

    has_parent = "parent" in cols
    has_level = "level" in cols
    rows: list[ImportRow] = []
    seen: dict[str, int] = {}
    level_stack: list[tuple[int, Optional[str]]] = []   # (livello, codice) delle righe precedenti

    for n, r in body:
        code = cell(r, "code") or None
        name = cell(r, "name")
        # senza intestazione e con una sola colonna: "01 Solai" → codice + nome
        if "code" not in cols and name and not code:
            m = re.match(r"^([\w.\-/]+)\s+(.+)$", name)
            if m and (m.group(1).isdigit() or CODE_SEPARATORS.search(m.group(1))):
                code, name = m.group(1), m.group(2)
        row = ImportRow(row=n, code=code, name=name, parent_code=None)
        if not name:
            row.error = "nome mancante" if code else "riga senza nome"
        elif code and code in seen:
            row.error = f"codice duplicato (riga {seen[code]})"
        if code and not row.error:
            seen[code] = n

        if row.error:
            rows.append(row)
            continue

        if has_parent:
            row.parent_code = cell(r, "parent") or None
        elif has_level:
            lv = cell(r, "level")
            try:
                level = int(float(lv)) if lv else 1
            except ValueError:
                row.error = f"livello non valido: {lv!r}"
                rows.append(row)
                continue
            while level_stack and level_stack[-1][0] >= level:
                level_stack.pop()
            row.parent_code = level_stack[-1][1] if level_stack else None
            if row.parent_code is None and level_stack:
                row.error = "livello sotto una voce senza codice"
            level_stack.append((level, code))
        elif code and CODE_SEPARATORS.search(code):
            # padre = codice senza l'ultimo segmento, risalendo finché ne trova uno nel file
            for m in reversed(list(CODE_SEPARATORS.finditer(code))):
                candidate = code[:m.start()]
                if candidate in seen:
                    row.parent_code = candidate
                    break
        rows.append(row)

    out.rows = rows
    return out
