# "><(((º> sabusabu <º)))><"
"""Report Excel della commessa: registro documenti, stato delle approvazioni, RFI. Ognuno contiene solo ciò
che chi esporta vede (RLS). Generazione sincrona con tetto di righe: per i volumi dell'MVP basta; se il test di
carico lo richiede si sposta nel worker con link a scadenza, senza cambiare l'URL di download."""
from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..deps import Ctx, ctx, require_project
from ..events import security_event
from .rfis import RFI_SQL
from .transmittals import STATE_LABEL

router = APIRouter(tags=["report"])

MAX_ROWS = 50_000
TZ = ZoneInfo("Europe/Rome")
HEAD_FILL = PatternFill("solid", fgColor="0C4577")
HEAD_FONT = Font(bold=True, color="FFFFFF")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

WF_STATUS = {"RUNNING": "In corso", "APPROVED": "Approvato", "REJECTED": "Rifiutato", "CANCELLED": "Annullato"}
RFI_STATUS = {"OPEN": "Aperta", "ANSWERED": "Con risposta", "CLOSED": "Chiusa"}
TITLES = {"documents": "Registro documenti", "approvals": "Stato approvazioni", "rfis": "RFI"}

Column = tuple[str, int]   # intestazione, larghezza


def _local(v):
    """Excel non conosce i fusi: le date si scrivono nell'ora italiana."""
    if isinstance(v, datetime):
        return v.astimezone(TZ).replace(tzinfo=None)
    return v


def _cell(ws, value, **style) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=_local(value))
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"                    # testo dell'utente, mai una formula
    elif isinstance(cell.value, datetime):
        cell.number_format = "dd/mm/yyyy hh:mm"
    for k, v in style.items():
        setattr(cell, k, v)
    return cell


def _sheet(wb: Workbook, title: str, columns: list[Column], rows: list[list]) -> None:
    """Foglio in streaming (write_only): tempo e memoria lineari nelle righe, anche a 50.000."""
    ws = wb.create_sheet(title)
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    ws.append([_cell(ws, c, fill=HEAD_FILL, font=HEAD_FONT, alignment=Alignment(vertical="center")) for c, _ in columns])
    for row in rows:
        ws.append([_cell(ws, v) for v in row])


def _info(wb: Workbook, project: dict, title: str, who: str, rows: int, truncated: bool) -> None:
    ws = wb.create_sheet("Info")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 70
    for label, value in [("Commessa", f"{project['code']} - {project['name']}"), ("Report", title),
                         ("Generato il", datetime.now(UTC)), ("Da", who), ("Righe", rows),
                         ("Nota", "Contiene solo le informazioni visibili a chi ha esportato."
                                  + (f" Troncato a {MAX_ROWS} righe." if truncated else ""))]:
        ws.append([_cell(ws, label, font=Font(bold=True)), _cell(ws, value)])


def _paths(conn, project_id: uuid.UUID) -> dict[uuid.UUID, str]:
    rows = {r["id"]: r for r in conn.execute(
        "SELECT id, parent_id, name FROM cde.containers WHERE project_id = %s", [project_id]).fetchall()}
    out: dict[uuid.UUID, str] = {}

    def path(cid):
        if cid not in out:
            r = rows.get(cid)
            out[cid] = "" if r is None else (f"{path(r['parent_id'])} / {r['name']}" if r["parent_id"] else r["name"])
        return out[cid]
    for cid in rows:
        path(cid)
    return out


def _documents(conn, project_id) -> tuple[list[Column], list[list]]:
    paths = _paths(conn, project_id)
    rows = conn.execute(
        """SELECT d.identifier, d.title, o.code AS originator, d.container_id, d.naming_fields ->> 'role' AS discipline,
                  lr.revision_code, lr.suitability, lr.state, lr.created_at, lr.original_filename, lr.size_bytes,
                  d.archived_at,
                  (SELECT r0.source_filename FROM cde.document_revisions r0      -- nome con cui era arrivato (rinomina)
                    WHERE r0.document_id = d.id AND r0.source_filename IS NOT NULL
                    ORDER BY r0.sequence DESC LIMIT 1) AS source_filename
           FROM cde.documents d JOIN cde.organizations o ON o.id = d.originator_organization_id
           JOIN LATERAL (SELECT r.revision_code, r.suitability, r.state, r.created_at, r.original_filename, f.size_bytes
                         FROM cde.document_revisions r JOIN cde.file_objects f ON f.id = r.file_object_id
                         WHERE r.document_id = d.id ORDER BY r.sequence DESC LIMIT 1) lr ON true
           WHERE d.project_id = %s ORDER BY d.identifier LIMIT %s""", [project_id, MAX_ROWS + 1]).fetchall()
    cols = [("Identificativo", 34), ("Titolo", 44), ("Originatore", 12), ("Cartella", 36), ("Disciplina", 11),
            ("Revisione", 10), ("Idoneità", 9), ("Stato", 12), ("Data revisione", 17), ("File", 36),
            ("Nome originale", 36), ("Dimensione (KB)", 15), ("Archiviato il", 17)]
    return cols, [[r["identifier"], r["title"], r["originator"], paths.get(r["container_id"], ""), r["discipline"],
                   r["revision_code"], r["suitability"], STATE_LABEL.get(r["state"], r["state"]), r["created_at"],
                   r["original_filename"], r["source_filename"], round(r["size_bytes"] / 1024, 1), r["archived_at"]]
                  for r in rows]


def _approvals(conn, project_id) -> tuple[list[Column], list[list]]:
    rows = conn.execute(
        """SELECT d.identifier, d.title, r.revision_code, t.name AS template_name, i.status, s.name AS step_name,
                  i.current_step, jsonb_array_length(t.definition -> 'steps') AS step_count, s.due_at,
                  (SELECT string_agg(u2.first_name || ' ' || u2.last_name, ', ' ORDER BY u2.last_name)
                     FROM cde.tasks k JOIN cde.users u2 ON u2.id = k.assignee_id
                     WHERE k.step_run_id = s.id AND k.status = 'OPEN') AS waiting_for,
                  u.first_name || ' ' || u.last_name AS started_by, i.started_at, i.finished_at,
                  (s.due_at IS NOT NULL AND s.due_at < now()) AS overdue
           FROM cde.workflow_instances i
           JOIN cde.workflow_templates t ON t.id = i.template_id
           JOIN cde.documents d ON d.id = i.document_id
           JOIN cde.document_revisions r ON r.id = i.revision_id
           LEFT JOIN cde.users u ON u.id = i.started_by
           LEFT JOIN cde.workflow_step_runs s ON s.instance_id = i.id AND s.step_index = i.current_step
                                              AND i.status = 'RUNNING'
           WHERE i.project_id = %s ORDER BY i.started_at DESC LIMIT %s""", [project_id, MAX_ROWS + 1]).fetchall()
    cols = [("Identificativo", 34), ("Titolo", 40), ("Revisione", 10), ("Workflow", 26), ("Stato", 12),
            ("Step corrente", 26), ("Avanzamento", 12), ("In attesa di", 32), ("Scadenza step", 17), ("Scaduto", 9),
            ("Avviato da", 22), ("Avviato il", 17), ("Concluso il", 17)]
    return cols, [[r["identifier"], r["title"], r["revision_code"], r["template_name"],
                   WF_STATUS.get(r["status"], r["status"]), r["step_name"],
                   f"{r['current_step'] + 1}/{r['step_count']}" if r["status"] == "RUNNING" else "",
                   r["waiting_for"], r["due_at"], "Sì" if r["overdue"] else "", r["started_by"], r["started_at"],
                   r["finished_at"]] for r in rows]


def _rfis(conn, project_id) -> tuple[list[Column], list[list]]:
    rows = conn.execute(
        "SELECT x.*, (SELECT string_agg(coalesce(o.code, u.first_name || ' ' || u.last_name), ', ')"
        "  FROM cde.rfi_recipients rr LEFT JOIN cde.organizations o ON o.id = rr.organization_id"
        "  LEFT JOIN cde.users u ON u.id = rr.user_id WHERE rr.rfi_id = x.id) AS recipients"
        " FROM (" + RFI_SQL + " WHERE f.project_id = %s) x ORDER BY x.number LIMIT %s",
        [project_id, MAX_ROWS + 1]).fetchall()
    cols = [("Codice", 11), ("Oggetto", 44), ("Stato", 13), ("Da", 26), ("Destinatari", 30), ("Aperta il", 17),
            ("Scadenza", 17), ("Scaduta", 9), ("Risposta il", 17), ("Chiusa il", 17), ("Messaggi", 10)]
    return cols, [[r["code"], r["subject"], RFI_STATUS.get(r["status"], r["status"]),
                   f"{r['author_org_code']} - {r['author_name']}", r["recipients"], r["created_at"], r["due_at"],
                   "Sì" if r["overdue"] else "", r["answered_at"], r["closed_at"], r["message_count"]] for r in rows]


BUILDERS = {"documents": _documents, "approvals": _approvals, "rfis": _rfis}


@router.get("/projects/{project_id}/reports/{kind}.xlsx")
def report(project_id: uuid.UUID, kind: Literal["documents", "approvals", "rfis"], c: Ctx = Depends(ctx)):
    with c.tx() as conn:
        project = require_project(conn, project_id)
        cols, rows = BUILDERS[kind](conn, project_id)
        who = conn.execute("SELECT first_name || ' ' || last_name AS n FROM cde.users WHERE id = cde.current_app_user()"
                           ).fetchone()["n"]
        truncated = len(rows) > MAX_ROWS
        rows = rows[:MAX_ROWS]
        security_event(conn, "REPORT_EXPORTED", user_id=c.user_id, project_id=project_id, ip=c.ip,
                       report=kind, rows=len(rows))
    wb = Workbook(write_only=True)
    _info(wb, project, TITLES[kind], who, len(rows), truncated)
    _sheet(wb, TITLES[kind], cols, rows)
    buf = io.BytesIO()
    wb.save(buf)
    name = f"{project['code']}-{kind}-{datetime.now(TZ):%Y%m%d}.xlsx"
    return Response(buf.getvalue(), media_type=XLSX, headers={
        "Content-Disposition": f'attachment; filename="{name}"', "Cache-Control": "no-store"})
