# "><(((º> sabusabu <º)))><"
"""Sprint 6 (G55): report Excel — registro documenti, approvazioni, RFI; solo ciò che chi esporta vede."""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from test_api_documents import PDF, des_name, upload
from test_api_states import as_user, transition, worker  # noqa: F401 - fixture riusate
from test_api_workflows import shared, template  # noqa: F401


def sheet(client, headers, project_id, kind):
    r = client.get(f"/projects/{project_id}/reports/{kind}.xlsx", headers=headers)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Info"
    return list(wb.worksheets[1].iter_rows(values_only=True))


def test_registro_documenti_rispetta_la_visibilita(client, api_world, as_user, worker):  # noqa: F811
    _, w = api_world
    name = des_name()
    upload(client, as_user("des_lead"), w.containers["des_arch"], name, PDF, title='=HYPERLINK("http://x")')
    worker()
    identifier = name.rsplit(".", 1)[0]
    rows = sheet(client, as_user("des_lead"), w.project_id, "documents")
    assert rows[0][0] == "Identificativo"
    mine = [r for r in rows if r[0] == identifier]
    assert mine and mine[0][7] == "WIP" and mine[0][1].startswith("=HYPERLINK")
    # WIP di un'altra organizzazione: non compare
    assert not [r for r in sheet(client, as_user("com_admin"), w.project_id, "documents") if r[0] == identifier]


def test_formule_non_eseguibili(client, api_world, as_user, worker):  # noqa: F811
    _, w = api_world
    upload(client, as_user("des_lead"), w.containers["des_arch"], des_name(), PDF, title="=1+1")
    worker()
    r = client.get(f"/projects/{w.project_id}/reports/documents.xlsx", headers=as_user("des_lead"))
    ws = load_workbook(io.BytesIO(r.content)).worksheets[1]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


def test_approvazioni_e_rfi(client, api_world, as_user, template, shared):  # noqa: F811
    _, w = api_world
    lead = as_user("des_lead")
    inst = client.post(f"/revisions/{shared['id']}/workflows", headers=lead, json={"template_id": template["id"]}).json()
    try:
        rows = sheet(client, as_user("com_admin"), w.project_id, "approvals")
        row = next(r for r in rows[1:] if r[2] == shared["revision_code"] and r[4] == "In corso" and r[3] == template["name"])
        assert row[5] == "Verifica interna" and row[6] == "1/3" and row[7]
    finally:
        client.post(f"/workflows/{inst['id']}/cancel", headers=lead, json={"reason": "fine test"})
    rfi = client.post(f"/projects/{w.project_id}/rfis", headers=lead, json={
        "subject": "Quota del solaio", "question": "Confermate +3,20?", "recipient_organization_ids": [str(w.orgs["COM"])]}).json()
    rows = sheet(client, lead, w.project_id, "rfis")
    assert any(r[0] == rfi["code"] and r[2] == "Aperta" and "COM" in r[4] for r in rows)
    assert not [r for r in sheet(client, as_user("sub1_user"), w.project_id, "rfis") if r[0] == rfi["code"]]


@pytest.mark.parametrize("kind", ["documents", "approvals", "rfis"])
def test_fuori_progetto(client, api_world, as_user, kind):  # noqa: F811
    _, w = api_world
    assert client.get(f"/projects/{w.project_id}/reports/{kind}.xlsx", headers=as_user("out_user")).status_code == 404
    assert client.get(f"/projects/{w.project_id}/reports/altro.xlsx", headers=as_user("des_lead")).status_code == 422
